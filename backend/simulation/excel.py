"""Excel export of the same structured result the API returns."""
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

SHEETS = [
    ("Resource Impact", "resource_impacts", [("Resource", "resource"), ("Location", "location"), ("Original Arrival", "original_arrival"), ("Delayed Arrival", "delayed_arrival"),
        ("Consumption Rate (/day)", "consumption_rate"), ("First Risk Date", "first_risk_date"), ("Minimum Projected", "minimum_projected"),
        ("Required Level", "required_level"), ("Deficit", "deficit"), ("Days Affected", "days_affected"), ("Unit", "unit"), ("Status", "status"), ("Reason", "reason")]),
    ("Expedition Impact", "expedition_impacts", [("Expedition ID", "expedition_id"), ("Expedition", "expedition"), ("Resource", "resource"), ("Location", "location"),
        ("Start", "start_date"), ("End", "end_date"), ("Required Qty", "required_quantity"), ("Status", "status"), ("Reason", "reason"), ("Recommended Action", "recommended_action")]),
    ("Alternative Actions", "alternatives", [("Expedition ID", "expedition_id"), ("Source", "source"), ("Resource", "resource"), ("Available", "available"),
        ("Required", "required"), ("Unit", "unit"), ("Feasible", "feasible"), ("Reason", "reason")]),
    ("Simulation Timeline", "timeline", [("Date", "date"), ("Resource", "resource"), ("Location", "location"), ("Opening", "opening"), ("Consumed", "consumed"),
        ("Incoming", "incoming"), ("Expedition Draw", "expedition_draw"), ("Available", "available"), ("Required", "required"), ("Status", "status")]),
]


def _cell(v):
    return str(v) if hasattr(v, "isoformat") else v


def _sheet(ws, cols, rows):
    ws.append([c[0] for c in cols])
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1C4E80")
    for r in rows:
        ws.append([_cell(r.get(k)) for _, k in cols])
    for i, (h, _) in enumerate(cols, 1):
        width = max([len(str(h))] + [len(str(_cell(r.get(cols[i - 1][1])) or "")) for r in rows[:200]])
        ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 70)
    ws.freeze_panes = "A2"


def build_workbook(result):
    wb = Workbook()
    ws = wb.active
    ws.title = "Impact Summary"
    s = result["scenario"]
    rows = [("Event", s["event"]), ("Cargo", "%s - %s %s %s" % (s["cargo_id"], s["quantity"], s["unit"], s["item"])), ("Destination", s["destination"]),
            ("Original Arrival", s["original_arrival"]), ("Delayed Arrival", s["delayed_arrival"]), ("Delay (days)", s["delay_days"]),
            ("Simulation Date", s["simulation_date"]), ("Overall Result", result["overall_status"]), ("Reason", result["overall_reason"]),
            ("Recommended Action", result["recommended_action"])] + [("Note", n) for n in result["notes"]]
    for k, v in rows:
        ws.append([k, _cell(v)])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
    ws.column_dimensions["A"].width = 22
    ws.column_dimensions["B"].width = 110
    for title, key, cols in SHEETS:
        _sheet(wb.create_sheet(title), cols, result[key])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
